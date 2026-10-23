import docx, io
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw

LOREM = ("The quick brown fox jumps over the lazy dog while the committee reviews the quarterly figures. "
         "Revenue grew steadily across all regions, although costs in logistics rose faster than expected. "
         "Management expects the trend to continue into the next financial year, subject to market conditions. ")

def para(d, text, size=11, bold=False, italic=False, align=None, color=None, font='Liberation Serif', space_after=6):
    p = d.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.name = font
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def rtl(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:bidi'); pPr.append(b)
    for r in p.runs:
        rPr = r._r.get_or_add_rPr(); x = OxmlElement('w:rtl'); rPr.append(x)

def picture(path, color):
    im = Image.new('RGB', (600, 300), color); dr = ImageDraw.Draw(im)
    for i in range(0, 600, 40): dr.line([(i, 0), (600 - i, 300)], fill='white', width=3)
    im.save(path)

# 1) English report: headings, styles, lists, table, image, header/footer page numbers
d = docx.Document()
sec = d.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7); sec.left_margin = sec.right_margin = Cm(2.5)
hp = sec.header.paragraphs[0]; hp.text = 'Quarterly Report — Confidential'; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '1'; r.append(t); fld.append(r); fp._p.append(fld)
d.add_heading('Annual Performance Review', 0)
d.add_heading('1. Overview', 1)
for i in range(3): para(d, LOREM * 3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
p = para(d, '', 11); r = p.add_run('Bold statement. '); r.bold = True; r = p.add_run('Italic remark. '); r.italic = True; r = p.add_run('Coloured warning.'); r.font.color.rgb = RGBColor(0xC0, 0, 0)
d.add_heading('2. Key points', 1)
for s in ['Sales up 12% year on year', 'Two new offices opened', 'Customer satisfaction at 91%']: d.add_paragraph(s, style='List Bullet')
for s in ['Review pricing', 'Hire analysts', 'Expand to new markets']: d.add_paragraph(s, style='List Number')
d.add_heading('3. Figures', 1)
tb = d.add_table(rows=4, cols=3); tb.style = 'Table Grid'
for i, row in enumerate([['Region', 'Revenue', 'Growth'], ['North', '1,200', '8%'], ['South', '980', '15%'], ['East', '1,450', '11%']]):
    for j, v in enumerate(row): tb.cell(i, j).text = v
para(d, '')
picture('chart.png', '#3366aa'); d.add_picture('chart.png', width=Inches(4))
para(d, 'Figure 1: Revenue by region.', 9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
d.add_heading('4. Outlook', 1)
for i in range(4): para(d, LOREM * 4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
d.save('en-report.docx')

# 2) Two-column newsletter
d = docx.Document()
d.add_heading('Community Newsletter', 0)
sectPr = d.sections[0]._sectPr; cols = OxmlElement('w:cols'); cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '720'); sectPr.append(cols)
for i in range(10): para(d, f'Story {i+1}. ' + LOREM * 2, 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
d.save('en-2col.docx')

# 3) Mixed Arabic / English, Arabic body with English terms and numbers
d = docx.Document()
for t in ['تقرير الأداء السنوي 2025', 'بلغت الإيرادات 1,200 مليون ريال بزيادة قدرها 12% عن العام الماضي، وفق تقرير شركة Acme Corp المنشور في يناير.',
          'وقد أشار المدير التنفيذي John Smith إلى أن خطة Vision 2030 أسهمت في نمو القطاع، مع توقعات باستمرار النمو خلال الربع الثالث (Q3).',
          'English paragraph inside an Arabic document, with numbers 3.14 and dates 2025-10-23.']:
    p = para(d, t, 14 if t.startswith('تقرير') else 12, bold=t.startswith('تقرير'), font='DejaVu Sans', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if any('؀' <= c <= 'ۿ' for c in t): rtl(p)
tb = d.add_table(rows=3, cols=3); tb.style = 'Table Grid'
for i, row in enumerate([['المنطقة', 'الإيرادات', 'النمو'], ['الرياض', '1,200', '8%'], ['جدة', '980', '15%']]):
    for j, v in enumerate(row): tb.cell(i, j).text = v
d.save('mixed-ar-en.docx')

# 4) Japanese
d = docx.Document()
for t in ['年次報告書', '本年度の売上高は前年比12％増加しました。物流費の上昇にもかかわらず、利益率は改善しました。来年度も成長が続く見込みです。' * 3]:
    para(d, t, 12, font='IPAGothic')
d.save('cjk.docx')

# 5) Landscape page
d = docx.Document()
from docx.enum.section import WD_ORIENT
s = d.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21)
d.add_heading('Landscape schedule', 1)
tb = d.add_table(rows=5, cols=6); tb.style = 'Table Grid'
for i in range(5):
    for j in range(6): tb.cell(i, j).text = f'R{i}C{j}' if i else f'Column {j+1}'
d.save('landscape.docx')
print('ok')
