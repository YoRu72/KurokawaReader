import os, zipfile, io
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PInches
from reportlab.lib.pagesizes import A5
from reportlab.pdfgen import canvas

OUT = os.path.join(os.path.dirname(__file__), '..', 'tests', 'fixtures')
os.makedirs(OUT, exist_ok=True)

def png(color, text):
    im = Image.new('RGB', (480, 300), color)
    d = ImageDraw.Draw(im)
    d.rectangle([20, 20, 460, 280], outline='white', width=6)
    d.text((40, 130), text, fill='white')
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()

def epub(path, title, author, lang, chapters, img_color):
    z = zipfile.ZipFile(path, 'w')
    z.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
    z.writestr('META-INF/container.xml', '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    manifest = ['<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>', '<item id="cover" href="images/cover.png" media-type="image/png" properties="cover-image"/>', '<item id="css" href="style.css" media-type="text/css"/>']
    spine = []
    nav = []
    z.writestr('OEBPS/images/cover.png', png(img_color, title))
    z.writestr('OEBPS/style.css', 'body{font-family:serif;color:#000;background:#fff} h1{color:darkred} .note{font-style:italic}')
    for i, (ctitle, paras) in enumerate(chapters):
        name = f'ch{i+1}.xhtml'
        z.writestr(f'OEBPS/images/fig{i+1}.png', png(img_color, f'{title} fig {i+1}'))
        body = f'<h1 id="t{i+1}">{ctitle}</h1>' + ''.join(f'<p>{p}</p>' for p in paras) + f'<figure><img src="images/fig{i+1}.png" alt="Figure {i+1} of {title}"/><figcaption>Figure {i+1}</figcaption></figure>' + (f'<p class="note">See <a href="ch{i+2}.xhtml#t{i+2}">the next chapter</a>.</p>' if i + 1 < len(chapters) else '')
        z.writestr(f'OEBPS/{name}', f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{lang}"><head><title>{ctitle}</title><link rel="stylesheet" href="style.css"/></head><body>{body}</body></html>')
        manifest.append(f'<item id="c{i+1}" href="{name}" media-type="application/xhtml+xml"/>')
        manifest.append(f'<item id="f{i+1}" href="images/fig{i+1}.png" media-type="image/png"/>')
        spine.append(f'<itemref idref="c{i+1}"/>')
        nav.append(f'<li><a href="{name}#t{i+1}">{ctitle}</a></li>')
    z.writestr('OEBPS/nav.xhtml', f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Nav</title></head><body><nav epub:type="toc"><ol>{"".join(nav)}</ol></nav></body></html>')
    z.writestr('OEBPS/content.opf', f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:{title}</dc:identifier><dc:title>{title}</dc:title><dc:creator>{author}</dc:creator><dc:language>{lang}</dc:language></metadata><manifest>{"".join(manifest)}</manifest><spine>{"".join(spine)}</spine></package>')
    z.close()

# Real books: the bundled starter EPUBs (public domain) replace any invented fixtures.
import shutil, glob
for f in glob.glob(os.path.join(os.path.dirname(__file__), '..', 'public', 'samples', '*.epub')):
    shutil.copy(f, OUT)

# PDF
c = canvas.Canvas(f'{OUT}/sample.pdf', pagesize=A5)
c.setTitle('Sample Report'); c.setAuthor('Report Writer')
for i in range(1, 9):
    c.setFont('Helvetica-Bold', 18); c.drawString(50, 540, f'Section {i}')
    c.setFont('Helvetica', 11)
    for j in range(20):
        c.drawString(50, 500 - j * 18, f'Page {i} line {j+1}: the quick brown fox jumps over the lazy dog.')
    c.bookmarkPage(f'p{i}'); c.addOutlineEntry(f'Section {i}', f'p{i}', level=0)
    c.showPage()
c.save()

# DOCX
d = Document()
d.core_properties.title = 'Meeting Notes'; d.core_properties.author = 'Secretary'
d.add_heading('Meeting Notes', 0)
d.add_heading('Agenda', 1)
for i in range(3): d.add_paragraph(f'Item {i+1}: discuss the roadmap and ' + 'details ' * 20, style='List Bullet')
d.add_heading('ملاحظات بالعربية', 1)
d.add_paragraph('هذه فقرةٌ عربيةٌ داخل مستند Word، تُختبر بها المحاذاة والخط.')
p = d.add_paragraph(); r = p.add_run('Bold text'); r.bold = True; p.add_run(' and '); r2 = p.add_run('italic text'); r2.italic = True
tbl = d.add_table(rows=3, cols=3); tbl.style = 'Table Grid'
for i in range(3):
    for j in range(3): tbl.cell(i, j).text = f'R{i+1}C{j+1}'
open('/tmp/fig.png', 'wb').write(png((60, 90, 60), 'docx image'))
d.add_picture('/tmp/fig.png', width=Inches(3))
d.save(f'{OUT}/notes.docx')

# XLSX
wb = Workbook(); ws = wb.active; ws.title = 'Budget'
ws.append(['Item', 'Qty', 'Unit price', 'Total', 'Date'])
import datetime
for i in range(1, 41):
    ws.append([f'Item {i}', i, 2.5 * i, f'=B{i+1}*C{i+1}', datetime.date(2026, 1, (i % 28) + 1)])
ws['A1'].font = ws['A1'].font.copy(bold=True)
ws2 = wb.create_sheet('ملخّص'); ws2.append(['البند', 'القيمة']); ws2.append(['الإجمالي', 1234.5])
wb.save(f'{OUT}/budget.xlsx')

# PPTX
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = 'Quarterly Review'; s.placeholders[1].text = 'Prepared by the team'
s2 = prs.slides.add_slide(prs.slide_layouts[1]); s2.shapes.title.text = 'Highlights'
tf = s2.placeholders[1].text_frame; tf.text = 'Revenue up 12%'
for t in ['New markets opened', 'Costs held flat']:
    para = tf.add_paragraph(); para.text = t; para.level = 1
s2.shapes.add_picture('/tmp/fig.png', PInches(1), PInches(4), width=PInches(3))
s2.notes_slide.notes_text_frame.text = 'Remember to thank the team.'
prs.save(f'{OUT}/review.pptx')

open(f'{OUT}/plain.txt', 'w', encoding='utf-8').write('A Plain Text Title\n\nFirst paragraph of the text file.\nStill the first paragraph.\n\nSecond paragraph here.\n\nفقرة عربية في ملف نصي.\n')
open(f'{OUT}/notes.md', 'w').write('# Markdown Notes\n\nSome **bold** text and a list:\n\n- one\n- two\n\n## Second heading\n\nMore text.\n')
open(f'{OUT}/data.csv', 'w').write('name,score\nAli,90\nBob,85\n')
print('fixtures written to', OUT)
